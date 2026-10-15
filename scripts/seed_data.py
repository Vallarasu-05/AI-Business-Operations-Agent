"""Generate PDF + DOCX copies of the sample .txt documents (to exercise the PDF/DOCX extractors).
Run from /backend:  python ../scripts/seed_data.py"""
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
from app.services.pdf_service import markdown_to_pdf  # noqa: E402
import docx  # noqa: E402

DOCS = ROOT / "data" / "sample_documents"
for txt in DOCS.glob("*.txt"):
    body = txt.read_text("utf-8")
    markdown_to_pdf("# " + txt.stem.replace("_", " ") + "\n\n" + body, str(DOCS / f"{txt.stem}.generated.pdf"))
    d = docx.Document()
    d.add_heading(txt.stem.replace("_", " "), 1)
    for line in body.splitlines():
        d.add_paragraph(line)
    d.save(str(DOCS / f"{txt.stem}.generated.docx"))
    print("generated", txt.stem)
print("Note: the mock Drive provider matches by filename prefix, so duplicates (txt+pdf+docx) will all be indexed.")
